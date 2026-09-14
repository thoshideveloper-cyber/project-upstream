"""Client-workbook (.xlsx) parsing — pure, no DB (WB-1, PHASE2_REFINEMENT_PLAN §2).

The three real client workbooks are not tidy tables: every sheet carries a *header
block* above the real column row (``"[Client] target/buyer name"``, ``"Exchange rate as
on date = X"``, a running count), the follow-up columns repeat their header four times,
and one workbook mixes master sheets, a shared email scheduler, research long-lists and
pivot tabs in one file. So this module does three things and only three:

1. **Locate** the real header row per sheet (never assume row 1) and read the header
   block above it (client label, exchange rate, count).
2. **Classify** the sheet — MASTER / SCHEDULE / CONTACTS / LONGLIST / IGNORE — from
   which column dictionary its header row matches.
3. **Map** cells to canonical fields using the §2.2 / §2.3 / §2.4 dictionaries verbatim,
   coercing types and *keeping the raw text* for anything ambiguous so the caller can
   flag it rather than guess.

Deciding what a row *means* (events, cadence, categories) is the importer's job, not
this module's — see ``services/workbook_import.py``.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

import openpyxl
from rapidfuzz import fuzz

# ── Sheet kinds ───────────────────────────────────────────────────────────────

MASTER = "MASTER"
SCHEDULE = "SCHEDULE"
CONTACTS = "CONTACTS"
LONGLIST = "LONGLIST"
IGNORE = "IGNORE"

# How far down to hunt for the real header row (the observed blocks are 2–7 rows).
_HEADER_SCAN_ROWS = 20
# A header row must match at least this many known column aliases to count.
_MIN_HEADER_HITS = 3
# rapidfuzz floor for a non-exact header→field guess.
_FUZZY_THRESHOLD = 88


def normalise_header(value: Any) -> str:
    """Lowercase, strip accents-of-punctuation, collapse whitespace.

    ``"Rev (INR Cr)"`` → ``"rev inr cr"``; ``"Email Id 1"`` → ``"email id 1"``.
    """
    text = str(value or "").replace("\xa0", " ")
    text = re.sub(r"[^0-9a-zA-Z]+", " ", text)
    return re.sub(r"\s+", " ", text).strip().lower()


def clean_text(value: Any) -> str | None:
    """Trim a cell to clean text, or None when it is blank."""
    if value is None:
        return None
    text = str(value).replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text or None


# ── Column dictionaries — PHASE2_REFINEMENT_PLAN §2.2 / §2.3 / §2.4 ───────────
# field → header aliases (normalised). Exact alias match wins; rapidfuzz is the
# fallback so a lightly-renamed column ("Website URL") still lands.

MASTER_COLUMNS: dict[str, list[str]] = {
    "company_name": ["company name", "company", "target name", "buyer name"],
    "hq": ["hq", "headquarters", "location"],
    "initial_email": ["initial email", "initial mail", "initial email date"],
    "status": ["status"],
    "category_raw": ["type", "company type", "counterparty type"],
    "rationale": ["rationale", "deal rationale"],
    "revenue_source": ["revenue from source", "revenue source", "revenue as sourced"],
    "revenue_inr_cr": ["rev inr cr", "revenue inr cr", "rev in inr cr", "revenue in inr cr"],
    "headcount": ["headcount", "employees", "employee count"],
    "website": ["website link", "website", "web link", "url"],
    "contact_name_1": ["contact name", "contact name 1", "contact person", "contact person 1"],
    "designation_1": ["designation", "designation 1", "title"],
    "email_1": ["email id", "email id 1", "email", "email 1"],
    "linkedin_1": ["linkedin", "linkedin 1", "linkedin url"],
    "contact_name_2": ["contact name 2", "contact person 2"],
    "designation_2": ["designation 2", "title 2"],
    "email_2": ["email id 2", "email 2"],
    "linkedin_2": ["linkedin 2", "linkedin url 2"],
    "relevant_investments": ["relevant investments", "relevant investment", "portfolio companies"],
    "bucket": ["bucket", "sourcing layer", "band"],
    "previous_contact": ["previous contact", "previously contacted"],
    "regarding": ["regarding"],
    "holding_company": ["holding company"],
    "holding_company_website": ["holding company website link", "holding company website"],
    "cheque_size": ["cheque size", "check size", "ticket size"],
    "notes": ["notes", "remarks", "comments"],
}

SCHEDULE_COLUMNS: dict[str, list[str]] = {
    "company_name": ["company name", "company"],
    "check_previous": [
        "check to previous sheets",
        "check to previous page",
        "check to previous sheets page",
        "check to previous",
    ],
    "regarding": ["regarding"],
    "initial_date": ["initial date", "initial email"],
    "status": ["status"],
    "follow_up": ["bi weekly follow up", "biweekly follow up", "follow up"],
    "days_block": ["follow ups in days", "follow ups days", "follow up in days"],
}

CONTACT_COLUMNS: dict[str, list[str]] = {
    "company": ["company", "company name"],
    "company_type": ["company type", "type"],
    "contact_person": ["contact person", "contact name", "name"],
    "designation": ["designation", "title"],
    "email_or_phone": ["email phone number", "email phone", "email", "phone number", "email id"],
    "reason": ["reason"],
    "engagement": ["engagement"],
    "date_connected": ["date connected", "connected on"],
    "mode": ["mode", "channel"],
    "poc": ["poc", "owner", "analyst"],
    "remark": ["remark", "sentiment", "outcome"],
    "comments": ["comments", "notes"],
}

# Columns the current data model has nowhere to put. Parsed and reported so the
# preview can say so out loud (never silently dropped) — the verbatim cell also
# survives on ``import_rows.raw``.
UNMAPPED_MASTER_FIELDS = (
    "holding_company",
    "holding_company_website",
    "previous_contact",
    "cheque_size",
)


# ── Header-row location + sheet classification ────────────────────────────────


def _match_field(header: str, dictionary: dict[str, list[str]]) -> str | None:
    """Resolve one normalised header to a canonical field: exact alias, then fuzzy."""
    if not header:
        return None
    for field_name, aliases in dictionary.items():
        if header in aliases:
            return field_name
    best_field, best_score = None, 0
    for field_name, aliases in dictionary.items():
        score = max((fuzz.ratio(header, alias) for alias in aliases), default=0)
        if score > best_score:
            best_field, best_score = field_name, score
    return best_field if best_score >= _FUZZY_THRESHOLD else None


def _row_headers(ws, row_index: int) -> list[str]:
    return [
        normalise_header(ws.cell(row=row_index, column=c).value)
        for c in range(1, (ws.max_column or 0) + 1)
    ]


def _score_row(headers: list[str], dictionary: dict[str, list[str]]) -> int:
    """How many *distinct* canonical fields this row's headers cover."""
    hits: set[str] = set()
    for header in headers:
        matched = _match_field(header, dictionary)
        if matched:
            hits.add(matched)
    return len(hits)


def _classify(headers: list[str]) -> tuple[str, int]:
    """Return (sheet kind, confidence score) for one candidate header row."""
    present = {h for h in headers if h}
    schedule_score = _score_row(headers, SCHEDULE_COLUMNS)
    contact_score = _score_row(headers, CONTACT_COLUMNS)
    master_score = _score_row(headers, MASTER_COLUMNS)

    has_followup = any(
        _match_field(h, SCHEDULE_COLUMNS) == "follow_up" for h in headers if h
    )
    has_company = any(
        _match_field(h, MASTER_COLUMNS) == "company_name" for h in headers if h
    )
    has_initial = any(
        _match_field(h, MASTER_COLUMNS) == "initial_email" for h in headers if h
    )
    has_person = "contact person" in present
    has_connected = "date connected" in present

    if has_company and has_followup:
        return SCHEDULE, schedule_score
    if has_person and has_connected:
        return CONTACTS, contact_score
    if has_company and has_initial:
        return MASTER, master_score
    if has_company and master_score >= _MIN_HEADER_HITS:
        # Company columns but no cadence anchor → a research long-list
        # (`PE names`, `Remaining PE companies`), not an active master sheet.
        return LONGLIST, master_score
    return IGNORE, 0


@dataclass
class SheetShape:
    """Everything the plan step needs to decide what to do with one tab."""

    title: str
    kind: str
    header_row: int | None
    headers: list[str]
    field_columns: dict[str, int | list[int]]
    data_row_count: int
    client_label: str | None = None
    exchange_rate: Decimal | None = None
    declared_count: int | None = None
    unmapped_headers: list[str] = field(default_factory=list)
    sample_rows: list[dict] = field(default_factory=list)
    regarding_values: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "title": self.title,
            "kind": self.kind,
            "header_row": self.header_row,
            "headers": self.headers,
            "data_row_count": self.data_row_count,
            "client_label": self.client_label,
            "exchange_rate": float(self.exchange_rate) if self.exchange_rate is not None else None,
            "declared_count": self.declared_count,
            "unmapped_headers": self.unmapped_headers,
            "sample_rows": self.sample_rows,
            "regarding_values": self.regarding_values,
        }


def _read_header_block(ws, header_row: int) -> tuple[str | None, Decimal | None, int | None]:
    """Read the free-form block above the column row: client label, FX rate, count."""
    client_label: str | None = None
    exchange_rate: Decimal | None = None
    declared_count: int | None = None
    for r in range(1, header_row):
        skip_col = -1
        for c in range(1, min(ws.max_column or 1, 8) + 1):
            if c == skip_col:
                continue  # the rate's value cell — not the running count
            value = ws.cell(row=r, column=c).value
            if value is None:
                continue
            text = clean_text(value)
            if text is None:
                continue
            low = text.lower()
            if "exchange rate" in low:
                # The rate sits in this cell ("… = 90.26") or the one beside it.
                inline = re.search(r"([0-9]+(?:\.[0-9]+)?)", text)
                neighbour = ws.cell(row=r, column=c + 1).value
                for candidate in (neighbour, inline.group(1) if inline else None):
                    parsed = as_decimal(candidate)
                    if parsed is not None and parsed > 0:
                        exchange_rate = parsed
                        break
                skip_col = c + 1
                continue
            if isinstance(value, (int, float)) and declared_count is None:
                declared_count = int(value)
                continue
            if client_label is None and len(text) > 3:
                client_label = text
    return client_label, exchange_rate, declared_count


def _last_data_row(ws, header_row: int, name_col: int) -> int:
    last = header_row
    for r in range(header_row + 1, (ws.max_row or header_row) + 1):
        if clean_text(ws.cell(row=r, column=name_col).value):
            last = r
    return last


def _resolve_columns(headers: list[str], dictionary: dict[str, list[str]]) -> dict[str, int]:
    """headers (0-based, normalised) → {field: 1-based column index}. First wins."""
    resolved: dict[str, int] = {}
    for i, header in enumerate(headers):
        matched = _match_field(header, dictionary)
        if matched and matched not in resolved:
            resolved[matched] = i + 1
    return resolved


def _shape_sheet(ws) -> SheetShape:
    """Locate the header row, classify the tab and resolve its columns."""
    best: tuple[int, str, int, list[str]] | None = None  # (score, kind, row, headers)
    for r in range(1, min(ws.max_row or 1, _HEADER_SCAN_ROWS) + 1):
        headers = _row_headers(ws, r)
        if sum(1 for h in headers if h) < _MIN_HEADER_HITS:
            continue
        kind, score = _classify(headers)
        if kind == IGNORE or score < _MIN_HEADER_HITS:
            continue
        if best is None or score > best[0]:
            best = (score, kind, r, headers)

    if best is None:
        return SheetShape(
            title=ws.title, kind=IGNORE, header_row=None, headers=[],
            field_columns={}, data_row_count=0,
        )

    _, kind, header_row, headers = best
    dictionary = {
        MASTER: MASTER_COLUMNS,
        LONGLIST: MASTER_COLUMNS,
        SCHEDULE: SCHEDULE_COLUMNS,
        CONTACTS: CONTACT_COLUMNS,
    }[kind]

    field_columns: dict[str, int | list[int]] = dict(_resolve_columns(headers, dictionary))
    if kind == SCHEDULE:
        # "Bi-weekly follow up" repeats ×4, and the trailing Done/days cells run from
        # the "Follow ups (in days)" header to the end — both are positional, not named.
        follow_up_cols = [
            i + 1 for i, h in enumerate(headers)
            if h and _match_field(h, SCHEDULE_COLUMNS) == "follow_up"
        ]
        field_columns["follow_up_cols"] = follow_up_cols
        days_col = field_columns.pop("days_block", None)
        field_columns.pop("follow_up", None)
        start = days_col if isinstance(days_col, int) else (
            (max(follow_up_cols) + 1) if follow_up_cols else len(headers) + 1
        )
        field_columns["done_cols"] = list(range(start, (ws.max_column or start) + 1))

    name_field = "company" if kind == CONTACTS else "company_name"
    name_col = field_columns.get(name_field)
    if not isinstance(name_col, int):
        return SheetShape(
            title=ws.title, kind=IGNORE, header_row=header_row, headers=headers,
            field_columns={}, data_row_count=0,
        )

    last_row = _last_data_row(ws, header_row, name_col)
    data_row_count = sum(
        1 for r in range(header_row + 1, last_row + 1)
        if clean_text(ws.cell(row=r, column=name_col).value)
    )
    client_label, exchange_rate, declared_count = _read_header_block(ws, header_row)

    mapped_cols = {v for v in field_columns.values() if isinstance(v, int)}
    for v in field_columns.values():
        if isinstance(v, list):
            mapped_cols.update(v)
    unmapped = [
        str(ws.cell(row=header_row, column=i + 1).value)
        for i, h in enumerate(headers)
        if h and (i + 1) not in mapped_cols
    ]

    regarding_values: list[str] = []
    reg_col = field_columns.get("regarding")
    if isinstance(reg_col, int):
        for r in range(header_row + 1, last_row + 1):
            text = clean_text(ws.cell(row=r, column=reg_col).value)
            if text and text not in regarding_values:
                regarding_values.append(text)

    return SheetShape(
        title=ws.title,
        kind=kind,
        header_row=header_row,
        headers=[
            str(ws.cell(row=header_row, column=i + 1).value or "")
            for i in range(len(headers))
        ],
        field_columns=field_columns,
        data_row_count=data_row_count,
        client_label=client_label,
        exchange_rate=exchange_rate,
        declared_count=declared_count,
        unmapped_headers=unmapped,
        regarding_values=regarding_values,
    )


# ── Value coercion ────────────────────────────────────────────────────────────

_NOT_A_VALUE = {"n/a", "na", "-", "--", "nil", "none", "tbd", "?"}
_DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%d %b %Y", "%d %B %Y")


def as_date(value: Any) -> date | None:
    """A real date, or None. Text tokens ("Got response") are never coerced."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def as_int(value: Any) -> int | None:
    text = clean_text(value)
    if not text or text.lower() in _NOT_A_VALUE:
        return None
    try:
        return int(float(text.replace(",", "")))
    except (TypeError, ValueError):
        return None


def as_decimal(value: Any) -> Decimal | None:
    text = clean_text(value)
    if not text or text.lower() in _NOT_A_VALUE:
        return None
    cleaned = re.sub(r"[,\s₹$€£]", "", text)
    try:
        return Decimal(cleaned)
    except (InvalidOperation, ValueError):
        return None


def as_revenue(value: Any) -> Decimal | None:
    """Revenue in INR Cr. A literal 0 is a spreadsheet formula artefact (the GAIL sheet
    has ``exchange rate = X``, so every converted cell evaluates to 0) — not a claim
    that the company earns nothing, so it reads as no-data."""
    parsed = as_decimal(value)
    if parsed is None or parsed == 0:
        return None
    return parsed


EMAIL_RE = re.compile(r"[^@\s]+@[^@\s]+\.[^@\s]+")


def split_email_phone(value: Any) -> tuple[str | None, str | None]:
    """The Contact List packs both into one 'Email / Phone number' column."""
    text = clean_text(value)
    if not text:
        return None, None
    match = EMAIL_RE.search(text)
    if match:
        return match.group(0), None
    if re.search(r"\d{5}", text):
        return None, text
    return None, None


# ── Row readers ───────────────────────────────────────────────────────────────


def _cell(ws, row: int, col: int | list[int] | None) -> Any:
    if not isinstance(col, int):
        return None
    return ws.cell(row=row, column=col).value


def read_rows(ws, shape: SheetShape) -> list[dict]:
    """Read a MASTER / LONGLIST / CONTACTS sheet into raw dicts keyed by canonical field.

    Values stay as text/date/number exactly as the cell holds them — interpretation
    (dates vs stop tokens, category aliases) belongs to the importer.
    """
    if shape.header_row is None or not shape.field_columns:
        return []
    name_field = "company" if shape.kind == CONTACTS else "company_name"
    name_col = shape.field_columns.get(name_field)
    if not isinstance(name_col, int):
        return []
    last_row = _last_data_row(ws, shape.header_row, name_col)

    rows: list[dict] = []
    for r in range(shape.header_row + 1, last_row + 1):
        if not clean_text(ws.cell(row=r, column=name_col).value):
            continue
        row: dict[str, Any] = {"_excel_row": r}
        for field_name, col in shape.field_columns.items():
            value = _cell(ws, r, col)
            if isinstance(value, datetime):
                row[field_name] = value.date().isoformat()
            elif isinstance(value, date):
                row[field_name] = value.isoformat()
            elif isinstance(value, Decimal):
                row[field_name] = str(value)
            else:
                row[field_name] = clean_text(value) if isinstance(value, str) else value
        rows.append(row)
    return rows


@dataclass
class ScheduleRow:
    """One Emailing-schedule row: the fixed anchor plus *evidence* of what was sent."""

    company_name: str
    regarding: str | None
    initial_date: date | None
    initial_token: str | None
    status_date: date | None
    status_token: str | None
    follow_up_dates: list[date | None]
    done_count: int
    check_previous: str | None


def read_schedule_rows(ws, shape: SheetShape) -> list[ScheduleRow]:
    """Read the Emailing-schedule tab.

    The trailing block is the sheet's own record of what actually went out: a ``Done``
    per completed follow-up, then a *number* = days until the next one. Counting the
    ``Done`` cells (not the dates, which are all four computed up front) is what tells
    us how many FOLLOW_UP events really happened.
    """
    if shape.header_row is None:
        return []
    name_col = shape.field_columns.get("company_name")
    if not isinstance(name_col, int):
        return []
    fu_cols = shape.field_columns.get("follow_up_cols") or []
    done_cols = shape.field_columns.get("done_cols") or []
    last_row = _last_data_row(ws, shape.header_row, name_col)

    out: list[ScheduleRow] = []
    for r in range(shape.header_row + 1, last_row + 1):
        name = clean_text(ws.cell(row=r, column=name_col).value)
        if not name:
            continue
        initial_raw = _cell(ws, r, shape.field_columns.get("initial_date"))
        status_raw = _cell(ws, r, shape.field_columns.get("status"))
        initial_date = as_date(initial_raw)
        status_date = as_date(status_raw)
        done = 0
        for col in done_cols:
            text = clean_text(ws.cell(row=r, column=col).value)
            if text and text.lower() == "done":
                done += 1
        out.append(
            ScheduleRow(
                company_name=name,
                regarding=clean_text(_cell(ws, r, shape.field_columns.get("regarding"))),
                initial_date=initial_date,
                initial_token=None if initial_date else clean_text(initial_raw),
                status_date=status_date,
                status_token=None if status_date else clean_text(status_raw),
                follow_up_dates=[as_date(ws.cell(row=r, column=c).value) for c in fu_cols],
                done_count=done,
                check_previous=clean_text(
                    _cell(ws, r, shape.field_columns.get("check_previous"))
                ),
            )
        )
    return out


# ── Public entry point ────────────────────────────────────────────────────────


@dataclass
class ParsedWorkbook:
    filename: str
    sheets: list[SheetShape]
    rows_by_sheet: dict[str, list[dict]]
    schedule_rows_by_sheet: dict[str, list[ScheduleRow]]

    def shape(self, title: str) -> SheetShape | None:
        return next((s for s in self.sheets if s.title == title), None)


def parse_workbook(raw: bytes, filename: str = "workbook.xlsx") -> ParsedWorkbook:
    """Open an .xlsx and shape every tab. Formula cells read their cached values."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    except Exception as exc:  # noqa: BLE001 — surfaced to the user as a 422
        raise ValueError(f"Could not read workbook: {exc}") from exc

    sheets: list[SheetShape] = []
    rows_by_sheet: dict[str, list[dict]] = {}
    schedule_rows: dict[str, list[ScheduleRow]] = {}
    try:
        for ws in wb.worksheets:
            shape = _shape_sheet(ws)
            sheets.append(shape)
            if shape.kind == SCHEDULE:
                schedule_rows[shape.title] = read_schedule_rows(ws, shape)
            elif shape.kind in (MASTER, CONTACTS, LONGLIST):
                rows = read_rows(ws, shape)
                rows_by_sheet[shape.title] = rows
                shape.sample_rows = rows[:5]
    finally:
        wb.close()

    return ParsedWorkbook(
        filename=filename,
        sheets=sheets,
        rows_by_sheet=rows_by_sheet,
        schedule_rows_by_sheet=schedule_rows,
    )
